"""jobs.xlsx: the day's CSVs as a sortable Excel workbook (excel.py, kept up to date by store.py)."""

import asyncio
import csv
import logging
import os
import time
from datetime import date, datetime, timedelta

import openpyxl
import pytest

import excel
import store as store_module
from decide import make_verdict
from store import MATCH_COLUMNS, REJECT_COLUMNS

JOB = {"job_id": "021026014440", "title": "ML Engineer", "company": "Acme", "location": "Pune",
       "experience_text": "", "link": "https://example.com/42"}
MATCH_HEADER = [column for column, _ in excel.MATCH_SHEET]
REJECT_HEADER = [column for column, _ in excel.REJECT_SHEET]


def today_folder(store):
    return store.output_dir / date.today().strftime("%Y-%m-%d")


def match(score: int, confidence: float) -> dict:
    return make_verdict(JOB, "jd", "pass", f"match {score}/10", confidence, match_score=score,
                        role_family="ai_ml_engineering", min_years=1, work_mode="remote")


def load(folder):
    return openpyxl.load_workbook(folder / "jobs.xlsx")


def data_rows(ws) -> list[dict]:
    header = [cell.value for cell in ws[1]]
    return [dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def write_csv(path, columns, rows) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=columns, restval="")
        writer.writeheader()
        writer.writerows(rows)


def test_workbook_is_created_after_a_match(store):
    store.append_match("naukri", JOB, match(8, 0.8), "mock")
    assert store.refresh_workbooks() is True

    wb = load(today_folder(store))  # loading at all = a valid workbook, even with an empty Rejected sheet
    assert wb.sheetnames == ["Matches", "Rejected"]

    matches, rejected = wb["Matches"], wb["Rejected"]
    assert [cell.value for cell in matches[1]] == MATCH_HEADER
    assert matches.tables["MatchesTable"].ref == "A1:L2"  # a real Excel table: header arrows to sort/filter
    assert matches.freeze_panes == "A2"

    assert [cell.value for cell in rejected[1]] == REJECT_HEADER
    assert rejected.max_row == 1 and not rejected.tables  # no rows yet: header only, no (invalid) empty table


def test_cells_are_typed(store):
    store.append_match("naukri", JOB, match(8, 0.8), "mock")
    store.refresh_workbooks()
    by_name = dict(zip(MATCH_HEADER, load(today_folder(store))["Matches"][2]))  # column name -> cell

    job_id = by_name["job_id"]
    assert job_id.value == "021026014440" and job_id.number_format == "@"  # text: leading zero kept
    assert by_name["match_score"].value == 8 and isinstance(by_name["match_score"].value, int)
    assert by_name["confidence"].value == 0.8 and by_name["confidence"].number_format == "0.00"
    assert by_name["min_years"].value == 1
    assert isinstance(by_name["time"].value, datetime) and by_name["time"].number_format == "yyyy-mm-dd hh:mm"
    assert by_name["link"].value == "open" and by_name["link"].hyperlink.target == "https://example.com/42"


def test_matches_sorted_by_score_then_confidence_across_sites(store):
    store.append_match("naukri", {**JOB, "job_id": "1"}, match(6, 0.9), "mock")
    store.append_match("glassdoor", {**JOB, "job_id": "2"}, match(10, 0.5), "mock")
    store.append_match("naukri", {**JOB, "job_id": "3"}, match(8, 0.7), "mock")
    store.append_match("indeed", {**JOB, "job_id": "4"}, match(8, 0.9), "mock")
    store.refresh_workbooks()

    rows = data_rows(load(today_folder(store))["Matches"])
    assert [(r["site"], r["match_score"], r["confidence"]) for r in rows] == [
        ("glassdoor", 10, 0.5), ("indeed", 8, 0.9), ("naukri", 8, 0.7), ("naukri", 6, 0.9),
    ]
    assert load(today_folder(store))["Matches"].tables["MatchesTable"].ref == "A1:L5"


def test_rejected_sheet_is_newest_first(tmp_path):
    write_csv(tmp_path / "naukri_rejected.csv", REJECT_COLUMNS, [
        {"time": "2026-10-15 09:00:00", "site": "naukri", "job_id": "1", "stage": "title", "reason": "a"},
        {"time": "2026-10-15 11:00:00", "site": "naukri", "job_id": "2", "stage": "jd", "reason": "b"},
    ])
    write_csv(tmp_path / "indeed_rejected.csv", REJECT_COLUMNS, [
        {"time": "2026-10-15 10:00:00", "site": "indeed", "job_id": "3", "stage": "title", "reason": "c"},
    ])
    excel.write_workbook(tmp_path)

    wb = load(tmp_path)
    assert [r["job_id"] for r in data_rows(wb["Rejected"])] == ["2", "3", "1"]
    assert wb["Rejected"].tables["RejectedTable"].ref == "A1:I4"
    assert wb["Matches"].max_row == 1  # no matches CSV at all: header only


def test_text_is_never_a_formula_and_control_characters_are_dropped(tmp_path):
    write_csv(tmp_path / "naukri_matches.csv", MATCH_COLUMNS, [
        {"time": "2026-10-15 09:00:00", "site": "naukri", "job_id": "7", "title": "=1+1", "company": "Bad\x0bCo",
         "match_score": "8", "min_years": "=1+1", "link": "javascript:alert(1)"},
    ])
    excel.write_workbook(tmp_path)

    cells = dict(zip(MATCH_HEADER, load(tmp_path)["Matches"][2]))
    assert (cells["title"].value, cells["title"].data_type) == ("=1+1", "s")
    assert (cells["min_years"].value, cells["min_years"].data_type) == ("=1+1", "s")  # not a number: text
    assert cells["company"].value == "BadCo"
    assert cells["link"].value == "javascript:alert(1)" and cells["link"].hyperlink is None  # not clickable


def test_characters_xml_forbids_are_dropped_everywhere(tmp_path):
    """U+FFFE / U+FFFF (e.g. from a "&#xFFFF;" on a job page): openpyxl saves them without complaint, and
    Excel then has to "repair" the file and drops the sheet's data."""
    write_csv(tmp_path / "naukri_matches.csv", MATCH_COLUMNS, [
        {"time": "2026-10-15 09:00:00￿", "site": "naukri", "job_id": "8", "title": "Odd￿Title",
         "company": "Odd￾Co", "match_score": "x￿", "link": "https://example.com/￿8"},
    ])
    excel.write_workbook(tmp_path)

    cells = dict(zip(MATCH_HEADER, load(tmp_path)["Matches"][2]))  # loading at all = well-formed XML
    assert (cells["title"].value, cells["company"].value) == ("OddTitle", "OddCo")
    assert cells["time"].value == datetime(2026, 10, 15, 9, 0)
    assert cells["match_score"].value == "x"  # not a number: kept as text
    assert cells["link"].hyperlink.target == "https://example.com/8"


def test_csv_formula_guard_is_not_shown_in_the_workbook(store):
    """store.py writes "'@Home" to the CSV so Excel won't run it. The workbook cell is text-only anyway,
    so it shows the real name, which you can then search and filter for."""
    job = {**JOB, "title": "- AI Engineer", "company": "@Home Solutions", "location": "+91 Pune",
           "link": '=HYPERLINK("http://x")'}
    store.append_match("naukri", job, match(8, 0.8), "mock")
    store.refresh_workbooks()

    cells = dict(zip(MATCH_HEADER, load(today_folder(store))["Matches"][2]))
    texts = [cells[c] for c in ("title", "company", "location", "link")]
    assert [cell.value for cell in texts] == ["- AI Engineer", "@Home Solutions", "+91 Pune", '=HYPERLINK("http://x")']
    assert all(cell.data_type == "s" for cell in texts)  # still text, never a formula


def test_appends_only_queue_the_workbook(store, monkeypatch):
    """A rebuild re-reads the whole day, so one per row would slow every screening down as the day grows.
    Rows that arrive together share one rebuild."""
    rebuilt = []
    monkeypatch.setattr(excel, "write_workbook", rebuilt.append)
    for job_id in ("1", "2", "3"):
        store.append_match("naukri", {**JOB, "job_id": job_id}, match(8, 0.8), "mock")
    store.append_rejected("indeed", JOB, make_verdict(JOB, "title", "reject", "x", 0.9), "mock")
    assert rebuilt == []

    store.refresh_workbooks()
    assert rebuilt == [today_folder(store)]  # one rebuild for all four rows
    store.refresh_workbooks()
    assert rebuilt == [today_folder(store)]  # nothing new: no rebuild


def test_locked_workbook_never_fails_and_is_retried(store, monkeypatch, caplog):
    """Excel open on jobs.xlsx: os.replace raises PermissionError on Windows. The CSV rows are safe, one
    warning is logged (not one per retry), no temp file is left behind, and the first refresh after Excel
    is closed brings the workbook up to date - without waiting for another row."""
    caplog.set_level(logging.INFO, logger="store")
    real_replace = os.replace

    def locked(src, dst):
        raise PermissionError(13, "Permission denied", str(dst))

    monkeypatch.setattr(os, "replace", locked)
    store.append_match("naukri", {**JOB, "job_id": "1"}, match(8, 0.8), "mock")
    assert store.refresh_workbooks() is False
    store.append_match("naukri", {**JOB, "job_id": "2"}, match(6, 0.8), "mock")
    assert store.refresh_workbooks() is False

    folder = today_folder(store)
    with open(folder / "naukri_matches.csv", encoding="utf-8-sig", newline="") as f:
        assert [row["job_id"] for row in csv.DictReader(f)] == ["1", "2"]
    assert not (folder / "jobs.xlsx").exists()
    assert not (folder / "jobs.xlsx.tmp").exists()
    warnings = [r.getMessage() for r in caplog.records if r.levelname == "WARNING"]
    assert warnings == [f"{folder.name}/jobs.xlsx is open in Excel - it is updated within "
                        f"{store_module.WORKBOOK_RETRY} s after you close it"]

    monkeypatch.setattr(os, "replace", real_replace)  # Excel closed
    assert store.refresh_workbooks() is True
    assert [r["job_id"] for r in data_rows(load(folder)["Matches"])] == ["1", "2"]
    assert "jobs.xlsx is up to date again" in caplog.messages


def test_startup_without_csvs_creates_nothing(store):
    store.queue_outdated_workbooks()  # what the server does at startup
    assert store.refresh_workbooks() is True
    assert not store.output_dir.exists()


def test_startup_rebuilds_outdated_workbooks(tmp_path, request):
    """Today: rows written just before a restart. Yesterday: the day ended while Excel had jobs.xlsx open.
    The day before: already up to date, so left alone."""
    output = tmp_path / "output"
    today, yesterday, before = (output / (date.today() - timedelta(days=n)).strftime("%Y-%m-%d") for n in range(3))
    for folder, job_id in ((today, "5"), (yesterday, "6"), (before, "7")):
        folder.mkdir(parents=True)
        write_csv(folder / "naukri_matches.csv", MATCH_COLUMNS, [
            {"time": "2026-10-15 09:00:00", "site": "naukri", "job_id": job_id, "title": "AI", "match_score": "6"},
        ])
    an_hour_ago = time.time() - 3600
    (yesterday / "jobs.xlsx").write_bytes(b"older than the CSV")
    os.utime(yesterday / "jobs.xlsx", (an_hour_ago, an_hour_ago))
    (before / "jobs.xlsx").write_bytes(b"newer than the CSV")
    os.utime(before / "naukri_matches.csv", (an_hour_ago, an_hour_ago))

    request.getfixturevalue("client")  # starts the server now, after the CSVs exist
    assert [r["job_id"] for r in data_rows(load(today)["Matches"])] == ["5"]
    assert [r["job_id"] for r in data_rows(load(yesterday)["Matches"])] == ["6"]
    assert (before / "jobs.xlsx").read_bytes() == b"newer than the CSV"


@pytest.fixture
def quick_workbooks(monkeypatch):
    monkeypatch.setattr(store_module, "WORKBOOK_DELAY", 0.01)


async def test_server_rebuilds_the_workbook_in_the_background(quick_workbooks, async_client, tmp_path):
    folder = tmp_path / "output" / date.today().strftime("%Y-%m-%d")
    reply = await async_client.post("/screen/title", json={"site": "naukri", "jobs": [
        {"job_id": "88", "title": "Python Trainer", "company": "EduCorp"},
    ]})
    assert reply.json()["results"][0]["decision"] == "reject"

    for _ in range(500):  # normally a few hundredths of a second
        if (folder / "jobs.xlsx").exists():
            break
        await asyncio.sleep(0.01)
    assert [r["job_id"] for r in data_rows(load(folder)["Rejected"])] == ["88"]


async def test_shutdown_saves_the_last_rows(app_env, tmp_path):
    """Rows from the last couple of seconds, not rebuilt yet, still reach jobs.xlsx when the server stops."""
    from app import app

    folder = tmp_path / "output" / date.today().strftime("%Y-%m-%d")
    async with app.router.lifespan_context(app):
        app.state.store.append_match("naukri", JOB, match(8, 0.8), "mock")
        assert not (folder / "jobs.xlsx").exists()  # WORKBOOK_DELAY not over yet
    assert [r["job_id"] for r in data_rows(load(folder)["Matches"])] == [JOB["job_id"]]
