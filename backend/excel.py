"""
jobs.xlsx: one day's CSV rows as an Excel workbook you can sort and filter.

The CSVs stay the source of truth (and they are what VS Code shows). Opened in Excel, a raw CSV gets
mangled: 12-digit job ids become 2.71E+11, Naukri ids lose their leading zero, the time column shows
########. So a couple of seconds after new CSV rows arrive, store.py rebuilds OUTPUT_DIR/YYYY-MM-DD/jobs.xlsx
from that day's CSVs, with real types:

    Matches    every <site>_matches.csv row, best match_score first (ties: highest position_score first)
    Rejected   every <site>_rejected.csv row, newest first

The CSVs are read by column NAME, including the <site>_<kind>.before-HHMMSS.csv files store.py set aside
when the columns changed: a column an older file doesn't have is just an empty cell.

Each sheet is an Excel Table: click a header arrow to sort or filter. The whole file is regenerated every
time, so edits saved in jobs.xlsx are lost; use Save As for a copy you want to keep notes in.
"""

import csv
import os
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

WORKBOOK_NAME = "jobs.xlsx"

# Characters XML 1.0 forbids. openpyxl's own ILLEGAL_CHARACTERS_RE misses U+FFFE / U+FFFF (and lone surrogates)
# and saves them as-is: Excel then has to "repair" the file and drops the sheet's data.
ILLEGAL_XML_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")

# (column, width in characters), in reading order - the CSVs keep their own order.
MATCH_SHEET = [
    ("match_score", 13), ("position_score", 15), ("skills_score", 13), ("confidence", 12), ("title", 45),
    ("company", 30), ("site", 11), ("location", 22), ("min_years", 11), ("work_mode", 12), ("role_family", 20),
    ("link", 45), ("time", 18), ("job_id", 16),
]
REJECT_SHEET = [
    ("time", 18), ("site", 11), ("stage", 8), ("title", 45), ("company", 30), ("reason", 40),
    ("confidence", 12), ("link", 45), ("job_id", 16),
]
NUMBER_COLUMNS = {"match_score", "position_score", "skills_score", "confidence", "min_years"}

# "naukri_matches.csv" and the set-aside "naukri_matches.before-134501.csv" (see store.py).
DAY_CSV_RE = re.compile(r"^(?P<site>[a-z]+)_(?P<kind>matches|rejected)(\.before-[\d-]+)?\.csv$")


def write_workbook(folder: Path) -> None:
    """Rebuild folder/jobs.xlsx from the CSVs in folder. Raises PermissionError while Excel has it open."""
    matches = _read_csvs(folder, "matches")
    matches.sort(key=lambda r: (_number(r.get("match_score")) or 0, _number(r.get("position_score")) or 0),
                 reverse=True)
    rejected = _read_csvs(folder, "rejected")
    rejected.sort(key=lambda r: r.get("time", ""), reverse=True)  # "YYYY-MM-DD HH:MM:SS" sorts as text

    wb = Workbook()
    _fill_sheet(wb.active, "Matches", MATCH_SHEET, matches)
    _fill_sheet(wb.create_sheet(), "Rejected", REJECT_SHEET, rejected)

    # Save to a temp file next to it, then swap it in: Excel never sees a half-written jobs.xlsx.
    # On Windows the swap fails while jobs.xlsx is open in Excel; the temp file is removed either way.
    # A fixed temp name is fine: store.py runs one rebuild at a time.
    tmp = folder / (WORKBOOK_NAME + ".tmp")
    try:
        wb.save(tmp)
        os.replace(tmp, folder / WORKBOOK_NAME)
    finally:
        tmp.unlink(missing_ok=True)


def day_csvs(folder: Path, kind: str) -> list[tuple[str, Path]]:
    """(site, path) of every <site>_<kind>*.csv in a day folder, old column layouts included."""
    found = []
    for path in sorted(folder.glob("*.csv")) if folder.exists() else []:
        m = DAY_CSV_RE.match(path.name)
        if m and m["kind"] == kind:
            found.append((m["site"], path))
    return found


def _read_csvs(folder: Path, kind: str) -> list[dict]:
    """All rows of all files, by column name (DictReader), so files with different headers mix fine."""
    rows = []
    for _, path in day_csvs(folder, kind):
        with open(path, encoding="utf-8-sig", newline="") as f:
            # key None = extra values in a broken row (more values than header columns): ignore them
            rows.extend({key: value or "" for key, value in row.items() if key is not None}
                        for row in csv.DictReader(f))
    return rows


def _fill_sheet(ws, name: str, columns: list[tuple[str, int]], rows: list[dict]) -> None:
    ws.title = name
    ws.append([column for column, _ in columns])
    for r, row in enumerate(rows, start=2):
        for c, (column, _) in enumerate(columns, start=1):
            _set_cell(ws.cell(row=r, column=c), column, row.get(column, ""))
    for c, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"  # the header stays visible while you scroll

    # Excel refuses a table without data rows, so an empty sheet keeps just the header.
    if rows:
        table = Table(displayName=f"{name}Table", ref=f"A1:{get_column_letter(len(columns))}{len(rows) + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)


def _set_cell(cell, column: str, text: str) -> None:
    """One CSV text value -> a typed Excel cell. Anything that isn't a valid number, time or web link
    becomes a text cell."""
    # Scraped text can hold characters XML forbids; drop them (in every branch, links included), not the workbook.
    text = ILLEGAL_XML_RE.sub("", text)
    if text == "":
        return  # e.g. min_years not stated: an empty cell, not the text ""
    if column in NUMBER_COLUMNS and _number(text) is not None:
        cell.value = _number(text)  # a real number, so sorting gives 10, 8, 6 (as text, "10" sorts before "6")
        if column == "confidence":
            cell.number_format = "0.00"
    elif column == "time" and _datetime(text) is not None:
        cell.value = _datetime(text)
        cell.number_format = "yyyy-mm-dd hh:mm"
    elif column == "link" and text.startswith(("http://", "https://")):
        # Only web links become clickable: the link comes from a job page, so never file:// or javascript:.
        # The full address as the text: some viewers (e.g. VS Code's spreadsheet viewers) drop the link when they
        # re-save the file; the address then stays visible and copyable instead of a dead "open".
        cell.value, cell.hyperlink, cell.style = text, text, "Hyperlink"
    else:
        if text[:1] == "'" and text[1:2] in ("=", "+", "-", "@"):
            # store.py's formula guard for the CSV ("'=..."). This cell is text-only, so it isn't needed here
            # and would show up as part of the name ("'@Home Solutions").
            text = text[1:]
        cell.value = text
        cell.data_type = "s"  # always text, never a formula - even "=..."
        if column == "job_id":
            cell.number_format = "@"  # text format: "021026014440" keeps its leading zero, no 2.71E+11


def _number(text: str | None) -> int | float | None:
    """CSV text -> number: "8" -> 8, "0.67" -> 0.67, "" or anything else -> None."""
    try:
        value = float(text)
    except (TypeError, ValueError):
        return None
    return int(value) if value.is_integer() else value


def _datetime(text: str) -> datetime | None:
    """CSV text in the format store.py writes -> datetime, anything else -> None (kept as text by _set_cell)."""
    try:
        return datetime.strptime(text, "%Y-%m-%d %H:%M:%S")
    except ValueError:
        return None
